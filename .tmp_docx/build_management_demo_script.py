from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pathlib import Path


OUT = Path(r"D:\Projects\ThucLuc_V2\.sources\artifacts\Kich_ban_demo_vai_tro_Quan_ly.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = "173F67"
BLUE = "EAF2FB"
PALE = "F5F8FC"
GOLD = "B68A2F"
LIGHT_GOLD = "F5EEDC"
GRAY = "5A6572"
BORDER = "D9D9D9"
WHITE = "FFFFFF"
BLACK = "000000"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=110, bottom=110, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_width(cell, cm):
    cell.width = Cm(cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, name="Arial", size=None, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_run_font(run, size=8.5, color=GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.7)
section.bottom_margin = Cm(1.55)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)
section.header_distance = Cm(0.7)
section.footer_distance = Cm(0.65)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
normal.font.size = Pt(10.2)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_after = Pt(4.5)
normal.paragraph_format.line_spacing = 1.12

title_style = styles["Title"]
title_style.font.name = "Arial"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
title_style.font.size = Pt(23)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor.from_string(BLACK)
title_style.paragraph_format.space_after = Pt(8)

for name, size, before, after in (("Heading 1", 16, 12, 6), ("Heading 2", 12.5, 9, 4), ("Heading 3", 10.8, 7, 3)):
    st = styles[name]
    st.font.name = "Arial"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(BLACK)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

footer = section.footer
fp = footer.paragraphs[0]
fp.text = "Kịch bản demo vai trò Quản lý  |  Hệ thống quản lý phát triển và ứng dụng CNTT"
fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
for r in fp.runs:
    set_run_font(r, size=8, color=GRAY)
add_page_number(footer.add_paragraph())


def add_label_value(label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(label + " ")
    set_run_font(r, bold=True, size=10.2)
    r = p.add_run(value)
    set_run_font(r, size=10.2)
    return p


def add_bullet(text, level=0, bold_prefix=None):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Cm(0.65 + level * 0.45)
    p.paragraph_format.first_line_indent = Cm(-0.25)
    p.paragraph_format.space_after = Pt(2.5)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_number(text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Cm(0.7)
    p.paragraph_format.first_line_indent = Cm(-0.32)
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run(text))
    return p


def add_two_col_table(rows, widths=(4.0, 12.7), header=None):
    extra = 1 if header else 0
    table = doc.add_table(rows=len(rows) + extra, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    idx = 0
    if header:
        for j, text in enumerate(header):
            cell = table.cell(0, j)
            set_cell_shading(cell, NAVY)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            set_run_font(run, bold=True, color=WHITE, size=9.6)
        set_repeat_table_header(table.rows[0])
        idx = 1
    for i, (left, right) in enumerate(rows, start=idx):
        row = table.rows[i]
        no_split(row)
        for j, value in enumerate((left, right)):
            cell = row.cells[j]
            set_cell_width(cell, widths[j])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if (i - idx) % 2 == 1:
                set_cell_shading(cell, PALE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            set_run_font(run, size=9.3, bold=(j == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_demo_table(rows):
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    headers = ["Thời lượng", "Thao tác trên màn hình", "Lời thoại đề xuất", "Kết quả cần chỉ ra"]
    widths = [1.65, 4.6, 7.55, 3.0]
    for i, text in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_width(cell, widths[i])
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell, top=100, bottom=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, bold=True, color=WHITE, size=8.7)
    set_repeat_table_header(table.rows[0])
    for ri, values in enumerate(rows):
        cells = table.add_row().cells
        no_split(table.rows[-1])
        for ci, value in enumerate(values):
            cell = cells[ci]
            set_cell_width(cell, widths[ci])
            set_cell_margins(cell, top=95, bottom=95)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri % 2 == 1:
                set_cell_shading(cell, PALE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ci == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(value)
            set_run_font(run, size=8.8, bold=(ci == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_script_section(number, title, time_text, objective, setup, rows, emphasis, transition, safe_note=None):
    doc.add_page_break()
    doc.add_heading(f"{number}  {title}", level=1)
    add_two_col_table([
        ("Thời lượng dự kiến", time_text),
        ("Mục tiêu trình bày", objective),
        ("Dữ liệu nên chuẩn bị", setup),
    ])
    doc.add_heading("Kịch bản thao tác và lời thoại", level=2)
    add_demo_table(rows)
    doc.add_heading("Điểm cần nhấn mạnh", level=2)
    for item in emphasis:
        add_bullet(item)
    if safe_note:
        p = doc.add_paragraph()
        r = p.add_run("Lưu ý khi demo  ")
        set_run_font(r, bold=True, color="9A4F12")
        r = p.add_run(safe_note)
        set_run_font(r)
    doc.add_heading("Câu chuyển sang màn hình tiếp theo", level=2)
    p = doc.add_paragraph()
    r = p.add_run(f'“{transition}”')
    set_run_font(r, italic=True, color=NAVY)


# Cover
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(38)
r = p.add_run("CÔNG AN NHÂN DÂN")
set_run_font(r, size=12, bold=True, color=NAVY)

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(36)
p.add_run("Kịch bản demo chức năng quản lý")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Hệ thống quản lý phát triển và ứng dụng công nghệ thông tin")
set_run_font(r, size=13, bold=True, color=NAVY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(12)
r = p.add_run("Vai trò trình diễn  Quản lý H05")
set_run_font(r, size=11.5, bold=True)

cover_table = doc.add_table(rows=4, cols=2)
cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER
cover_table.autofit = False
cover_table.style = "Table Grid"
for i, (a, b) in enumerate([
    ("Tài khoản dự kiến", "h05.user"),
    ("Thời lượng", "25 phút trình bày và 5 phút hỏi đáp"),
    ("Ngày trình bày", "07/10/2026"),
    ("Phạm vi", "Quản lý quy trình báo cáo và theo dõi dữ liệu các đơn vị"),
]):
    set_cell_width(cover_table.cell(i, 0), 4.2)
    set_cell_width(cover_table.cell(i, 1), 9.6)
    set_cell_shading(cover_table.cell(i, 0), LIGHT_GOLD)
    for j, txt in enumerate((a, b)):
        cell = cover_table.cell(i, j)
        set_cell_margins(cell, top=130, bottom=130)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p2 = cell.paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        set_run_font(p2.add_run(txt), size=10, bold=(j == 0))
set_table_borders(cover_table, color="D1C19A")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(28)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Mục tiêu")
set_run_font(r, bold=True, size=11.5)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.left_indent = Cm(2.0)
p.paragraph_format.right_indent = Cm(2.0)
r = p.add_run("Trình bày một luồng quản lý khép kín từ thiết lập đợt báo cáo, theo dõi tiến độ, xử lý yêu cầu bổ sung đến tra cứu kết quả, xem hiện trạng và nhận diện biến động.")
set_run_font(r, size=10.5)

doc.add_page_break()
doc.add_heading("1  Phạm vi demo và thông điệp chính", level=1)
p = doc.add_paragraph("Tài khoản Quản lý dùng để điều hành quy trình báo cáo và theo dõi dữ liệu của các đơn vị thuộc phạm vi quản lý. Tài khoản này không trực tiếp nhập dữ liệu nghiệp vụ tại từng module. Trong buổi demo, người trình bày cần giữ mạch nội dung theo ba câu hỏi quản lý dưới đây.")
for r in p.runs: set_run_font(r)
add_two_col_table([
    ("Câu hỏi 1", "Đợt báo cáo đang được tổ chức như thế nào và các đơn vị đã thực hiện đến đâu"),
    ("Câu hỏi 2", "Khi số liệu đã nộp cần điều chỉnh thì ai yêu cầu, ai phê duyệt và hệ thống kiểm soát ra sao"),
    ("Câu hỏi 3", "Tại một đơn vị, dữ liệu hiện tại là gì và đã thay đổi thế nào so với kỳ trước"),
], header=("Câu hỏi quản lý", "Chức năng trả lời"))

doc.add_heading("Thông điệp mở đầu có thể nói", level=2)
p = doc.add_paragraph()
r = p.add_run("“Tài khoản tôi đang sử dụng là vai trò Quản lý. Phạm vi của vai trò này là tổ chức kỳ báo cáo, theo dõi tiến độ các đơn vị, xử lý nhu cầu bổ sung và khai thác dữ liệu đã nộp. Các màn hình nhập liệu nghiệp vụ không xuất hiện vì việc nhập thuộc trách nhiệm của từng đơn vị.”")
set_run_font(r, italic=True, color=NAVY)

doc.add_heading("Các chức năng xuất hiện trên menu", level=2)
add_two_col_table([
    ("Yêu cầu bổ sung", "Xem xét, duyệt hoặc từ chối yêu cầu chỉnh sửa số liệu đã nộp"),
    ("Tiến độ báo cáo", "Theo dõi tỷ lệ và trạng thái nộp theo từng đơn vị"),
    ("Kỳ báo cáo", "Tạo và điều hành vòng đời của một kỳ thu thập"),
    ("Mẫu báo cáo", "Quản lý cấu trúc module và tần suất của biểu mẫu"),
    ("Tra cứu báo cáo", "Tìm báo cáo đã nộp, xem trước và xuất tệp"),
    ("Hiện trạng và biến động", "Xem dữ liệu hiện tại và so sánh chỉ số giữa các kỳ"),
], header=("Chức năng", "Ý nghĩa đối với quản lý"))

doc.add_heading("Nguyên tắc trình bày", level=2)
for text in [
    "Mỗi màn hình bắt đầu bằng câu hỏi quản lý mà màn hình đó trả lời.",
    "Nêu số liệu đang thấy trên màn hình trước khi bấm vào chi tiết.",
    "Chỉ thực hiện thao tác làm thay đổi dữ liệu trên kỳ và đơn vị đã chuẩn bị cho demo.",
    "Nếu không có bản ghi phù hợp, chuyển sang dữ liệu dự phòng thay vì tạo vội dữ liệu trong lúc trình bày.",
]: add_bullet(text)

doc.add_page_break()
doc.add_heading("2  Chuẩn bị trước buổi trình bày", level=1)
doc.add_heading("Danh sách kiểm tra trước 30 phút", level=2)
check_rows = [
    ("□", "Đăng nhập bằng h05.user và xác nhận góc dưới menu hiển thị H05 User"),
    ("□", "Kiểm tra sáu chức năng trong nhóm Báo cáo đều mở được, không xuất hiện thông báo lỗi"),
    ("□", "Chọn sẵn một kỳ đang mở có nhiều trạng thái đơn vị để demo Tiến độ báo cáo"),
    ("□", "Chuẩn bị một yêu cầu bổ sung ở trạng thái Chờ duyệt trên kỳ demo"),
    ("□", "Chuẩn bị một mẫu báo cáo có nhiều module và một mẫu chưa gắn kỳ nếu cần minh họa xóa"),
    ("□", "Chuẩn bị ít nhất một báo cáo đã nộp có thể xem trước PDF và tải Excel"),
    ("□", "Chuẩn bị hai kỳ cùng mẫu có dữ liệu để xem Biến động"),
    ("□", "Chuẩn bị một đơn vị có dữ liệu chi tiết ở module Dự án CNTT hoặc Nhân lực CNTT"),
    ("□", "Mở sẵn thư mục tải xuống và đóng các cửa sổ có thông tin không liên quan"),
    ("□", "Đặt trình duyệt ở mức thu phóng 90 hoặc 100 phần trăm và tắt thông báo hệ điều hành"),
]
table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
set_table_borders(table)
for i, h in enumerate(("Đã kiểm tra", "Nội dung")):
    set_cell_width(table.cell(0, i), (2.0, 14.7)[i])
    set_cell_shading(table.cell(0, i), NAVY)
    p = table.cell(0, i).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(p.add_run(h), bold=True, color=WHITE, size=9.5)
set_repeat_table_header(table.rows[0])
for ri, (mark, text) in enumerate(check_rows):
    cells = table.add_row().cells
    no_split(table.rows[-1])
    for i, val in enumerate((mark, text)):
        set_cell_width(cells[i], (2.0, 14.7)[i])
        set_cell_margins(cells[i])
        if ri % 2: set_cell_shading(cells[i], PALE)
        p = cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
        set_run_font(p.add_run(val), size=9.3, bold=(i == 0))

doc.add_heading("Dữ liệu demo đề xuất", level=2)
add_two_col_table([
    ("Kỳ chính", "[DEMO] Báo cáo đầy đủ Quý 3 năm 2026 hoặc kỳ đang mở có dữ liệu tương đương"),
    ("Đơn vị chính", "Một Công an tỉnh hoặc thành phố có dữ liệu đã nộp"),
    ("Module chi tiết", "Dự án CNTT để xem số dự án và tổng kinh phí; Nhân lực CNTT để xem số lượng nhân sự"),
    ("Cặp kỳ so sánh", "Hai kỳ dùng cùng một mẫu, có ít nhất một chỉ số tăng hoặc giảm"),
    ("Yêu cầu bổ sung", "Một yêu cầu Chờ duyệt với lý do rõ ràng và hạn bổ sung hợp lý"),
])

doc.add_heading("Hai chế độ demo", level=2)
add_two_col_table([
    ("An toàn", "Mở hộp thoại và giải thích thao tác, sau đó bấm Hủy. Dùng khi đang trình bày trên dữ liệu thật."),
    ("Đầy đủ", "Thực hiện Duyệt, Từ chối, Tạo kỳ hoặc Thêm mẫu trên dữ liệu demo đã chuẩn bị. Chỉ dùng khi được phép thay đổi dữ liệu."),
], header=("Chế độ", "Cách thực hiện"))

add_script_section(
    "3", "Mở đầu từ Trang chủ", "1 phút",
    "Xác nhận đúng vai trò, đúng phạm vi và giới thiệu mạch demo.",
    "Đăng nhập sẵn bằng h05.user và đứng tại Trang chủ.",
    [
        ("0:20", "Chỉ vào tên tài khoản H05 User ở góc dưới bên trái.", "Tôi đang đăng nhập bằng vai trò Quản lý H05. Vai trò này quản lý quy trình báo cáo và khai thác dữ liệu đã nộp của các đơn vị.", "Khán giả nhận biết đúng vai trò đang trình diễn."),
        ("0:20", "Mở nhóm Báo cáo trên menu.", "Các chức năng được tổ chức theo vòng đời: thiết lập mẫu và kỳ, theo dõi tiến độ, xử lý bổ sung, sau đó tra cứu và phân tích dữ liệu.", "Menu thể hiện đúng phạm vi quản lý."),
        ("0:20", "Di chuột qua các nhóm nghiệp vụ nhưng không mở.", "Tài khoản quản lý không trực tiếp vào các màn nhập liệu nghiệp vụ. Việc nhập thuộc tài khoản đơn vị; quản lý theo dõi qua báo cáo đã chốt.", "Làm rõ nguyên tắc phân quyền."),
    ],
    [
        "Quyền truy cập và menu phải đồng nhất: chức năng có quyền mới xuất hiện và mới truy cập được.",
        "Không mở đầu bằng chi tiết kỹ thuật; tập trung vào trách nhiệm quản lý và câu hỏi nghiệp vụ.",
    ],
    "Đầu tiên, tôi xin trình bày cách thiết lập khung thu thập dữ liệu thông qua Mẫu báo cáo và Kỳ báo cáo.",
)

add_script_section(
    "4", "Mẫu báo cáo", "3 phút",
    "Giải thích cách quy định nội dung thu thập cho từng đợt báo cáo.",
    "Có sẵn một mẫu nhiều module; nếu demo tạo mới, dùng mã DEMO_2026 và không lưu trên dữ liệu thật.",
    [
        ("0:35", "Chọn Mẫu báo cáo. Chỉ vào các cột Mã mẫu, Tên mẫu, Tần suất, Module và Trạng thái.", "Mẫu báo cáo định nghĩa một đợt sẽ thu những nhóm dữ liệu nào và theo tần suất nào. Nhờ đó từng kỳ không phải cấu hình lại từ đầu.", "Danh sách mẫu và trạng thái hoạt động."),
        ("0:35", "Nhập từ khóa hoặc chọn bộ lọc Tần suất và Trạng thái.", "Quản lý có thể tìm nhanh theo mã, tên hoặc module; đồng thời lọc theo tần suất và trạng thái sử dụng.", "Bộ lọc thu hẹp đúng danh sách."),
        ("1:10", "Bấm Thêm mẫu. Chỉ vào Mã mẫu, Tên mẫu, Tần suất, Mô tả và Danh sách module.", "Khi tạo mẫu, tôi xác định mã quản lý, tên hiển thị, chu kỳ thu thập và các module cần báo cáo. Danh sách module quyết định phạm vi dữ liệu của kỳ sử dụng mẫu này.", "Hộp thoại thể hiện cấu trúc mẫu."),
        ("0:25", "Mở một mẫu hiện có bằng biểu tượng bút chì.", "Mẫu có thể cập nhật khi chưa bị ràng buộc bởi quy trình đang vận hành. Hệ thống kiểm soát thao tác theo quyền Thêm, Sửa và Xóa.", "Nút thao tác xuất hiện theo quyền."),
        ("0:15", "Bấm Hủy để đóng hộp thoại.", "Trong buổi demo này tôi không lưu thay đổi để giữ nguyên dữ liệu chuẩn bị sẵn.", "Trở về danh sách an toàn."),
    ],
    [
        "Mẫu báo cáo là cấu hình nghiệp vụ, không phải dữ liệu do đơn vị nhập.",
        "Một mẫu gom nhiều module và có thể được tái sử dụng cho nhiều kỳ cùng cấu trúc.",
        "Mẫu đã gắn với kỳ sẽ bị hạn chế xóa để bảo toàn lịch sử.",
    ],
    "Sau khi đã có cấu trúc thu thập, bước tiếp theo là mở một Kỳ báo cáo cụ thể với thời gian và trạng thái rõ ràng.",
    "Không bấm Lưu mẫu nếu đang đứng trên dữ liệu thật. Nếu cần demo lưu, dùng mẫu DEMO riêng và xóa sau buổi trình bày khi mẫu chưa được sử dụng."
)

add_script_section(
    "5", "Kỳ báo cáo", "4 phút",
    "Trình bày cách tạo đợt thu thập và kiểm soát vòng đời của kỳ.",
    "Một kỳ đang mở, một kỳ đã đóng và một kỳ đã khóa; chuẩn bị sẵn dữ liệu để minh họa trạng thái.",
    [
        ("0:35", "Chọn Kỳ báo cáo và chỉ vào các bộ lọc trạng thái, năm và ô tìm kiếm.", "Màn hình này là nơi quản lý toàn bộ các đợt thu thập. Tôi có thể lọc theo trạng thái, theo năm hoặc tìm trực tiếp theo tên kỳ.", "Danh sách kỳ và trạng thái hiện tại."),
        ("1:00", "Bấm Tạo kỳ mới. Chỉ vào mẫu báo cáo, tên kỳ, năm, quý hoặc tháng và thời gian hiệu lực.", "Một kỳ luôn gắn với một mẫu báo cáo. Khi tạo kỳ, quản lý xác định phạm vi nội dung, chu kỳ và thời gian các đơn vị được phép nhập.", "Hộp thoại tạo kỳ đầy đủ trường nghiệp vụ."),
        ("0:30", "Chỉ vào nút Tạo và mở ngay rồi bấm Hủy.", "Tạo và mở ngay giúp rút gọn thao tác khi cấu hình đã hoàn tất. Trong demo an toàn tôi đóng hộp thoại mà không tạo dữ liệu mới.", "Không phát sinh kỳ ngoài kế hoạch."),
        ("1:00", "Chỉ vào một kỳ Đang mở, sau đó một kỳ Đã đóng và nút Mở lại.", "Kỳ đang mở cho phép đơn vị thực hiện báo cáo. Đóng kỳ dừng tiếp nhận thông thường; khi có lý do phù hợp, quản lý có thể mở lại kỳ đã đóng.", "Thể hiện các chuyển trạng thái hợp lệ."),
        ("0:40", "Chỉ vào thao tác Khóa kỳ và cảnh báo xác nhận.", "Khóa là trạng thái kết thúc vĩnh viễn của kỳ. Hệ thống cảnh báo rõ vì kỳ đã khóa không thể mở lại. Tôi chỉ thực hiện thao tác này sau khi đã đối chiếu tiến độ và kết quả nộp.", "Khán giả hiểu khác biệt giữa Đóng và Khóa."),
        ("0:15", "Bấm biểu tượng Xem tiến độ nộp trên một kỳ.", "Từ kỳ báo cáo, quản lý có thể đi thẳng sang màn theo dõi tiến độ của đúng kỳ đã chọn.", "Chuyển đúng màn Tiến độ báo cáo."),
    ],
    [
        "Chuẩn bị, Đang mở, Đã đóng và Khóa là các mốc kiểm soát của quy trình.",
        "Đóng có thể mở lại; Khóa không thể mở lại.",
        "Không dùng thao tác Khóa để minh họa trực tiếp trên kỳ nghiệp vụ đang vận hành.",
    ],
    "Khi kỳ đã mở, câu hỏi quan trọng nhất của quản lý là đơn vị nào đã nộp, đơn vị nào đang nhập và đơn vị nào cần đôn đốc.",
    "Nếu chỉ cần giới thiệu, dừng ở hộp xác nhận và bấm Hủy. Không khóa thử một kỳ thật."
)

add_script_section(
    "6", "Tiến độ báo cáo", "4 phút",
    "Cho thấy khả năng giám sát toàn bộ đơn vị và xử lý ngay trường hợp cần bổ sung.",
    "Chọn kỳ có đồng thời đơn vị Đã nộp, Đang nhập và Chưa nộp; có ít nhất một đơn vị Đã nộp để mở yêu cầu bổ sung.",
    [
        ("0:35", "Chọn Tiến độ báo cáo và chọn kỳ trong danh sách.", "Tôi chọn một kỳ để xem tiến độ trên toàn bộ các đơn vị địa phương thuộc phạm vi quản lý.", "Kỳ, trạng thái và khoảng thời gian hiển thị."),
        ("0:45", "Chỉ vào thanh tiến độ và sáu ô thống kê.", "Phần tổng quan cho biết tỷ lệ đã nộp và phân nhóm rõ: đã nộp, đã xác nhận, đang bổ sung, đang nhập và chưa nộp. Quản lý nhìn một lần là biết điểm nghẽn nằm ở đâu.", "Số tổng và từng trạng thái khớp danh sách."),
        ("0:45", "Cuộn bảng đơn vị và chỉ vào tên, trạng thái, thời gian nộp.", "Bảng chi tiết giúp xác định đơn vị cụ thể, trạng thái hiện tại và thời điểm nộp gần nhất để phục vụ đôn đốc hoặc đối chiếu.", "Trạng thái từng đơn vị rõ ràng."),
        ("1:15", "Tại một đơn vị Đã nộp, bấm Tạo yêu cầu. Nhập lý do và chọn hạn bổ sung.", "Nếu phát hiện số liệu cần chỉnh, quản lý có thể tạo yêu cầu trực tiếp tại đây. Tôi ghi rõ nội dung cần bổ sung và thời hạn hoàn thành; hệ thống sẽ mở lại báo cáo của đúng đơn vị và đúng kỳ.", "Hộp thoại gắn đúng đơn vị và kỳ."),
        ("0:40", "Bấm Hủy trong chế độ an toàn hoặc Gửi yêu cầu trên dữ liệu demo.", "Sau khi gửi, trạng thái đơn vị chuyển sang Đang bổ sung và yêu cầu được theo dõi ở chức năng Yêu cầu bổ sung.", "Luồng từ giám sát sang xử lý được liên kết."),
    ],
    [
        "Quản lý theo dõi theo từng kỳ, không cần mở từng module nhập liệu.",
        "Mỗi trạng thái có ý nghĩa hành động khác nhau: đôn đốc, đối chiếu hoặc yêu cầu bổ sung.",
        "Yêu cầu bổ sung được gắn với đúng kỳ và đúng đơn vị.",
    ],
    "Tôi chuyển sang Yêu cầu bổ sung để trình bày cách quản lý kiểm soát việc mở lại số liệu đã nộp.",
    "Chỉ bấm Gửi yêu cầu nếu đang sử dụng đơn vị và kỳ demo. Nếu không, mở hộp thoại, giải thích rồi bấm Hủy."
)

add_script_section(
    "7", "Yêu cầu bổ sung", "3 phút",
    "Trình bày cơ chế kiểm soát việc chỉnh sửa báo cáo đã nộp.",
    "Có một yêu cầu Chờ duyệt và một yêu cầu đã xử lý để minh họa lịch sử trạng thái.",
    [
        ("0:35", "Chọn Yêu cầu bổ sung. Chọn kỳ và lọc trạng thái Chờ duyệt.", "Ở vai trò Quản lý, màn hình hiển thị yêu cầu từ các đơn vị. Tôi có thể lọc theo kỳ, đơn vị và trạng thái để ưu tiên các yêu cầu đang chờ xử lý.", "Danh sách chỉ còn yêu cầu Chờ duyệt."),
        ("0:35", "Chỉ vào lý do, hạn bổ sung và thời điểm gửi của một yêu cầu.", "Quản lý có đủ thông tin để đánh giá: đơn vị đề nghị gì, lý do vì sao, gửi khi nào và dự kiến hoàn thành lúc nào.", "Thông tin quyết định được thể hiện trên một dòng."),
        ("0:50", "Bấm Duyệt và chỉ vào hạn bổ sung trong hộp thoại.", "Khi duyệt, hệ thống mở lại snapshot để đơn vị chỉnh sửa và nộp lại. Hạn bổ sung có thể được xác nhận hoặc điều chỉnh trước khi phê duyệt.", "Nút Xác nhận duyệt và hạn bổ sung."),
        ("0:35", "Đóng hộp Duyệt, sau đó bấm Từ chối trên yêu cầu demo khác.", "Nếu đề nghị không phù hợp, quản lý từ chối và bắt buộc ghi rõ lý do. Nội dung này giúp đơn vị hiểu quyết định và tạo dấu vết xử lý.", "Lý do từ chối là trường bắt buộc."),
        ("0:25", "Bấm Hủy để không thay đổi dữ liệu.", "Toàn bộ trạng thái từ Chờ duyệt đến Đã duyệt, Từ chối, Đang bổ sung và Hoàn thành được lưu để theo dõi xuyên suốt.", "Không phát sinh thay đổi ngoài kịch bản."),
    ],
    [
        "Báo cáo đã nộp không được sửa tự do; phải đi qua yêu cầu và quyết định của quản lý.",
        "Duyệt yêu cầu sẽ mở lại đúng snapshot; Từ chối phải có lý do.",
        "Trạng thái giúp theo dõi trách nhiệm và tiến độ xử lý.",
    ],
    "Sau khi các đơn vị nộp hoặc nộp lại, quản lý cần tra cứu bản đã chốt và xuất đúng biểu mẫu phục vụ tổng hợp.",
    "Không duyệt thử trên yêu cầu thật nếu đơn vị chưa được thông báo. Dùng yêu cầu demo hoặc chỉ mở hộp thoại."
)

add_script_section(
    "8", "Tra cứu báo cáo", "4 phút",
    "Minh họa cách tìm, kiểm tra và xuất báo cáo đã nộp của nhiều đơn vị.",
    "Có ít nhất một báo cáo đã nộp hoặc đã khóa, PDF tạo được và trình duyệt cho phép tải tệp.",
    [
        ("0:40", "Chọn Tra cứu báo cáo. Giữ chế độ Báo cáo gần nhất từng đơn vị.", "Chế độ mặc định cho tôi một bản báo cáo mới nhất của mỗi đơn vị, phù hợp khi cần tổng hợp nhanh tình hình hiện tại.", "Mỗi đơn vị xuất hiện với phiên bản gần nhất."),
        ("0:45", "Chuyển sang Theo một kỳ báo cáo cụ thể; chọn kỳ và trạng thái.", "Khi cần đối chiếu theo một đợt, tôi chọn kỳ cụ thể và có thể lọc tiếp theo trạng thái báo cáo.", "Danh sách thay đổi theo kỳ đã chọn."),
        ("0:30", "Nhập tên một đơn vị và bật Chỉ phiên bản mới nhất.", "Tìm kiếm theo đơn vị và tùy chọn phiên bản giúp truy xuất đúng hồ sơ cần kiểm tra mà không phải dò toàn bộ danh sách.", "Danh sách còn đúng đơn vị mục tiêu."),
        ("0:55", "Bấm biểu tượng con mắt để Xem trước PDF.", "Bản xem trước phản ánh biểu mẫu đã chốt tại thời điểm nộp. Quản lý có thể kiểm tra bố cục và nội dung trước khi tải xuống.", "PDF mở trong hộp thoại xem trước."),
        ("0:45", "Đóng xem trước; chỉ vào nút tải Excel và PDF.", "Cùng một snapshot có thể xuất Excel để xử lý số liệu hoặc PDF để lưu trữ, trình ký và chia sẻ theo biểu mẫu thống nhất.", "Hai định dạng xuất rõ ràng."),
        ("0:25", "Nếu báo cáo tổng hợp, bấm Chi tiết đóng góp theo đơn vị.", "Với báo cáo tổng hợp, hệ thống cho biết từng đơn vị đã đóng góp dữ liệu gì theo module, giúp truy vết nguồn của số liệu tổng.", "Danh sách đơn vị và chi tiết module."),
    ],
    [
        "Tra cứu đọc dữ liệu đã nộp, không đọc trực tiếp dữ liệu đang nhập.",
        "Phiên bản và trạng thái giúp phân biệt bản hiện hành với bản đã được thay thế.",
        "Excel phục vụ khai thác số liệu; PDF phục vụ trình bày và lưu hồ sơ.",
    ],
    "Tra cứu cho biết hồ sơ đã nộp. Phần cuối tôi trình bày cách biến các hồ sơ đó thành góc nhìn quản lý về hiện trạng và xu hướng thay đổi.",
    "Thử mở trước PDF trước giờ trình bày để tránh chờ dịch vụ sinh tệp. Nếu mạng chậm, ưu tiên bảng dữ liệu và nút tải Excel."
)

add_script_section(
    "9", "Hiện trạng và biến động", "5 phút",
    "Trả lời hai câu hỏi quản lý: dữ liệu hiện tại của đơn vị là gì và đã thay đổi ra sao giữa hai kỳ.",
    "Có một kỳ chứa đủ module; có hai kỳ cùng mẫu để so sánh; chọn sẵn đơn vị có biến động rõ.",
    [
        ("0:40", "Chọn Hiện trạng và biến động, giữ thẻ Hiện trạng và chọn kỳ báo cáo.", "Hiện trạng trả lời câu hỏi: tại kỳ đã chọn, đơn vị đang có dữ liệu gì ở từng nghiệp vụ. Đây là màn khai thác, không phải màn nhập liệu.", "Danh sách đơn vị và các thẻ module xuất hiện."),
        ("0:40", "Tìm theo tên hoặc mã đơn vị rồi chọn một đơn vị.", "Tài khoản quản lý có thể tìm nhanh trong phạm vi các đơn vị cấp tỉnh được quản lý. Khi chọn đơn vị, toàn bộ module có dữ liệu được hiển thị ở bên phải.", "Đơn vị mục tiêu được tô chọn."),
        ("0:50", "Bấm module Dự án CNTT hoặc Nhân lực CNTT.", "Mỗi thẻ cho một chỉ số chính để đọc nhanh. Khi mở rộng, hệ thống hiển thị các chỉ số chi tiết của module, ví dụ số dự án, tổng kinh phí hoặc tổng nhân sự CNTT.", "Bảng chỉ số hiện tại xuất hiện."),
        ("0:40", "Bấm Danh sách chi tiết và tìm một bản ghi.", "Nếu cần đi từ chỉ số tổng về dữ liệu cụ thể, quản lý mở Danh sách chi tiết. Nội dung này được lấy từ dữ liệu đã chốt của kỳ báo cáo.", "Danh sách bản ghi đúng module và đơn vị."),
        ("0:40", "Đóng chi tiết, chuyển sang thẻ Biến động.", "Biến động dùng hai kỳ cùng một mẫu để so sánh, bảo đảm các chỉ số có cùng ý nghĩa nghiệp vụ.", "Bộ chọn Kỳ đích và Kỳ đối chiếu xuất hiện."),
        ("0:55", "Chọn Kỳ đích, Kỳ đối chiếu, Module hoặc Tất cả module.", "Tôi chọn kỳ cần đánh giá và kỳ làm mốc. Có thể xem toàn bộ module hoặc tập trung vào một nghiệp vụ cụ thể.", "Các ô tổng hợp Đang cải thiện, Cần chú ý, Ổn định và Thiếu dữ liệu."),
        ("0:35", "Chọn đơn vị có cảnh báo; mở một module có chỉ số giảm.", "Hệ thống hiển thị giá trị kỳ đối chiếu, kỳ đích, chênh lệch và đánh giá. Ví dụ nhân sự từ 5 còn 3 được đánh dấu Cần chú ý; tỷ lệ đạt tăng được đánh dấu Cải thiện.", "Người xem hiểu cả giá trị và chiều biến động."),
    ],
    [
        "Hiện trạng trả lời dữ liệu đang là bao nhiêu; Biến động trả lời đã thay đổi thế nào.",
        "Tài khoản quản lý xem được nhiều đơn vị nhưng không cần quyền vào màn nhập của từng module.",
        "Hai kỳ so sánh phải dùng cùng mẫu để các chỉ số tương thích.",
        "Giảm không luôn đồng nghĩa với xấu; hệ thống đánh giá theo ý nghĩa của từng chỉ số.",
        "Thiếu dữ liệu đối chiếu được tách riêng để quản lý không nhầm với trạng thái ổn định.",
    ],
    "Như vậy, vai trò Quản lý có thể đi từ tổ chức đợt báo cáo đến theo dõi thực hiện, kiểm soát điều chỉnh và cuối cùng khai thác dữ liệu phục vụ điều hành.",
    "Nếu cặp kỳ không có dữ liệu tương thích, chuyển ngay sang cặp kỳ dự phòng đã kiểm tra trước."
)

doc.add_page_break()
doc.add_heading("10  Kết thúc phần trình bày", level=1)
doc.add_heading("Lời kết đề xuất", level=2)
p = doc.add_paragraph()
r = p.add_run("“Qua sáu chức năng vừa trình bày, vai trò Quản lý kiểm soát được toàn bộ vòng đời báo cáo. Mẫu và kỳ xác định phạm vi thu thập; Tiến độ và Yêu cầu bổ sung kiểm soát quá trình thực hiện; Tra cứu lưu giữ kết quả đã nộp; Hiện trạng và biến động chuyển dữ liệu thành thông tin phục vụ theo dõi, đánh giá và chỉ đạo.”")
set_run_font(r, italic=True, color=NAVY)

doc.add_heading("Ba thông điệp cần để người nghe nhớ", level=2)
for text in [
    "Quản lý điều hành và giám sát; đơn vị chịu trách nhiệm nhập và nộp dữ liệu.",
    "Mọi chỉnh sửa sau khi nộp đều đi qua luồng yêu cầu bổ sung có trạng thái và dấu vết xử lý.",
    "Dữ liệu đã chốt có thể tra cứu, xuất tệp, xem hiện trạng và so sánh biến động theo đơn vị và module.",
]: add_number(text)

doc.add_heading("Cách mời câu hỏi", level=2)
p = doc.add_paragraph()
r = p.add_run("“Tôi xin kết thúc phần demo. Tôi sẵn sàng làm rõ thêm về phạm vi dữ liệu của vai trò Quản lý, quy trình mở lại báo cáo hoặc cách đọc các chỉ số hiện trạng và biến động.”")
set_run_font(r, italic=True, color=NAVY)

doc.add_heading("Phân bổ thời gian tổng thể", level=2)
add_two_col_table([
    ("Mở đầu và phạm vi", "1 phút"),
    ("Mẫu báo cáo", "3 phút"),
    ("Kỳ báo cáo", "4 phút"),
    ("Tiến độ báo cáo", "4 phút"),
    ("Yêu cầu bổ sung", "3 phút"),
    ("Tra cứu báo cáo", "4 phút"),
    ("Hiện trạng và biến động", "5 phút"),
    ("Kết luận", "1 phút"),
], header=("Nội dung", "Thời lượng"))

doc.add_page_break()
doc.add_heading("11  Câu hỏi dự kiến và câu trả lời ngắn", level=1)
qa = [
    ("Vì sao tài khoản quản lý không vào từng module nhập liệu", "Vì nhập liệu thuộc trách nhiệm của đơn vị. Quản lý khai thác dữ liệu đã nộp qua Tra cứu, Hiện trạng và Biến động; cách tách này bảo đảm đúng trách nhiệm và hạn chế sửa dữ liệu nguồn."),
    ("Quản lý có xem được tất cả đơn vị không", "Tài khoản quản lý xem các đơn vị cấp tỉnh trong phạm vi được cấu hình. Phạm vi được kiểm soát ở backend, không chỉ ẩn hiện trên giao diện."),
    ("Đóng kỳ và khóa kỳ khác nhau thế nào", "Đóng kỳ dừng tiếp nhận nhưng còn có thể mở lại khi cần. Khóa kỳ là kết thúc vĩnh viễn và không thể mở lại."),
    ("Đơn vị đã nộp có tự sửa được không", "Không. Đơn vị phải có yêu cầu bổ sung được xử lý, hoặc quản lý chủ động tạo yêu cầu từ màn Tiến độ. Sau đó hệ thống mới mở lại đúng báo cáo để bổ sung."),
    ("Có biết đơn vị nào chưa nộp không", "Có. Màn Tiến độ hiển thị số tổng, tỷ lệ, nhóm trạng thái và danh sách cụ thể từng đơn vị."),
    ("Có tải được báo cáo không", "Có. Tra cứu báo cáo hỗ trợ xem trước PDF, tải PDF và xuất Excel từ snapshot đã nộp."),
    ("Số liệu hiện tại của một nghiệp vụ tại một đơn vị xem ở đâu", "Vào Hiện trạng và biến động, chọn thẻ Hiện trạng, chọn kỳ, tìm đơn vị rồi mở module cần xem. Có thể mở tiếp Danh sách chi tiết."),
    ("Biến động được tính như thế nào", "Hệ thống đối chiếu chỉ số đã chốt của hai kỳ dùng cùng mẫu, hiển thị giá trị hai kỳ, chênh lệch và đánh giá theo chiều ý nghĩa của từng chỉ số."),
    ("Nếu kỳ trước không có dữ liệu thì sao", "Hệ thống đánh dấu Thiếu dữ liệu đối chiếu thay vì kết luận tăng, giảm hoặc ổn định."),
    ("Có thể xem lịch sử phiên bản không", "Có. Tra cứu theo kỳ có thể hiển thị phiên bản và trạng thái; tùy chọn Chỉ phiên bản mới nhất giúp chuyển giữa góc nhìn hiện hành và lịch sử."),
    ("Mẫu báo cáo đã sử dụng có xóa được không", "Mẫu đang được gắn với kỳ sẽ không cho xóa để bảo toàn báo cáo và lịch sử. Có thể ngừng sử dụng theo chính sách quản trị dữ liệu."),
    ("Dữ liệu so sánh có lấy từ màn nhập hiện tại không", "Không. Phân tích sử dụng chỉ số được chốt cùng snapshot tại thời điểm nộp, nhờ đó kết quả không thay đổi theo dữ liệu đang chỉnh sửa."),
]
add_two_col_table(qa, widths=(5.5, 11.2), header=("Câu hỏi", "Câu trả lời đề xuất"))

doc.add_page_break()
doc.add_heading("12  Phương án xử lý tình huống trong lúc demo", level=1)
recovery = [
    ("Danh sách không có dữ liệu", "Kiểm tra lại kỳ đã chọn và bộ lọc; bấm Xóa lọc hoặc chuyển sang kỳ dự phòng đã kiểm tra."),
    ("PDF mở chậm", "Đóng hộp xem trước, tiếp tục trình bày bảng và nút xuất Excel; nói rõ PDF được sinh từ snapshot đã chốt."),
    ("Không thấy nút thao tác", "Giải thích nút được hiển thị theo quyền hành động. Không đổi tài khoản hoặc phân quyền ngay trong buổi demo."),
    ("Không có yêu cầu Chờ duyệt", "Dùng một bản ghi đã xử lý để giải thích trạng thái; mở màn Tiến độ và minh họa hộp Tạo yêu cầu nhưng không gửi."),
    ("Không có cặp kỳ so sánh", "Chuyển về Hiện trạng để trình bày dữ liệu hiện tại; nêu điều kiện hai kỳ phải cùng mẫu và dùng ảnh hoặc số liệu dự phòng nếu có."),
    ("Thông báo lỗi hệ thống", "Không thử lặp lại nhiều lần. Ghi nhận tên màn hình, thao tác và thời điểm; chuyển sang chức năng kế tiếp để giữ mạch trình bày."),
    ("Bấm nhầm thao tác thay đổi dữ liệu", "Dừng ở hộp xác nhận và chọn Hủy. Không xác nhận Khóa kỳ, Duyệt hoặc Xóa nếu không nằm trong dữ liệu demo."),
    ("Màn hình quá dài", "Dùng tìm kiếm hoặc bộ lọc thay vì cuộn nhanh; giữ con trỏ tại khu vực đang thuyết minh."),
]
add_two_col_table(recovery, widths=(4.6, 12.1), header=("Tình huống", "Cách xử lý"))

doc.add_heading("Các thao tác không nên thực hiện trên dữ liệu thật", level=2)
for text in [
    "Khóa kỳ chỉ để minh họa bằng hộp xác nhận, không bấm xác nhận cuối.",
    "Không xóa hoặc ngừng mẫu đang được sử dụng.",
    "Không duyệt hoặc từ chối yêu cầu của đơn vị thật nếu chưa thống nhất trước.",
    "Không tạo nhiều bản ghi thử trong lúc trình bày.",
    "Không thay đổi phân quyền để xử lý nhanh một nút đang thiếu.",
]: add_bullet(text)

doc.add_heading("Ghi chú cá nhân trước khi lên trình bày", level=2)
for label in ["Kỳ chính", "Kỳ dự phòng", "Đơn vị chính", "Đơn vị dự phòng", "Yêu cầu Chờ duyệt", "Module có biến động rõ"]:
    p = doc.add_paragraph()
    r = p.add_run(f"{label}: ")
    set_run_font(r, bold=True)
    r = p.add_run("........................................................................................................")
    set_run_font(r, color=GRAY)

# Document properties
doc.core_properties.title = "Kịch bản demo chức năng quản lý"
doc.core_properties.subject = "Hướng dẫn thao tác và lời thoại trình bày vai trò Quản lý H05"
doc.core_properties.author = "Cục Công nghệ thông tin"
doc.core_properties.keywords = "demo, quản lý, báo cáo, tiến độ, hiện trạng, biến động"

doc.save(OUT)
print(OUT)
